"""Tests for DocQA loaders + the docs_root sandbox.

Real fixtures per format in tmp_path (docx via python-docx; a minimal valid PDF built
inline; txt/md/html written directly). Focus: containment (no escape) + typed errors
(never a raw exception). Run in conda `local-ai-agent-env-1`: `pytest`.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

from local_ai_agent.modules.docqa.loaders import (
    DocAccessError,
    DocLoadError,
    DocNotFound,
    DocTooLarge,
    UnsupportedFormat,
    is_supported,
    list_documents,
    load_bytes,
    load_document,
    resolve_within,
)


# --------------------------------------------------------------------------- #
# fixtures / helpers
# --------------------------------------------------------------------------- #
def _make_pdf(text: str) -> bytes:
    """A minimal single-page PDF with one text run + a correct xref → pypdf extracts `text`."""
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
    ]
    stream = b"BT /F1 24 Tf 72 700 Td (" + text.encode("latin-1") + b") Tj ET"
    objs.append(b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream")
    objs.append(b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj" + body + b"endobj\n"
    xref_pos = len(out)
    n = len(objs) + 1
    out += b"xref\n0 " + str(n).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += b"trailer<</Size " + str(n).encode() + b"/Root 1 0 R>>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    return out


@pytest.fixture
def root(tmp_path: Path) -> Path:
    (tmp_path / "a.txt").write_text("hello txt", encoding="utf-8")
    (tmp_path / "b.md").write_text("# title\n\nbody **md**", encoding="utf-8")
    (tmp_path / "c.html").write_text(
        "<html><head><style>x{}</style></head><body><h1>Hi</h1>"
        "<script>evil()</script><p>world</p></body></html>", encoding="utf-8")
    (tmp_path / "d.pdf").write_bytes(_make_pdf("Hello PDF"))
    import docx
    doc = docx.Document()
    doc.add_paragraph("docx line one")
    doc.add_paragraph("docx line two")
    doc.save(str(tmp_path / "e.docx"))
    (tmp_path / "ignore.bin").write_bytes(b"\x00\x01\x02") # unsupported
    sub = tmp_path / "sub"
    sub.mkdir()
    (sub / "nested.txt").write_text("nested", encoding="utf-8")
    (tmp_path / ".hidden.txt").write_text("secret", encoding="utf-8")
    return tmp_path


# --------------------------------------------------------------------------- #
# NORMAL — each format loads
# --------------------------------------------------------------------------- #
def test_load_txt(root):
    assert load_document(root, "a.txt") == "hello txt"


def test_load_md(root):
    assert "body **md**" in load_document(root, "b.md")


def test_load_html_strips_tags_and_scripts(root):
    text = load_document(root, "c.html")
    assert "Hi" in text and "world" in text
    assert "evil" not in text and "<h1>" not in text and "x{}" not in text


def test_load_pdf(root):
    assert "Hello PDF" in load_document(root, "d.pdf")


def test_load_docx(root):
    text = load_document(root, "e.docx")
    assert "docx line one" in text and "docx line two" in text


def test_list_documents_sorted_relative_supported_only(root):
    docs = list_documents(root)
    assert docs == ["a.txt", "b.md", "c.html", "d.pdf", "e.docx", "sub/nested.txt"] or \
           sorted(docs) == docs # sorted, relative
    assert "ignore.bin" not in docs # unsupported excluded
    assert ".hidden.txt" not in docs # hidden excluded
    assert "sub/nested.txt" in docs # recursive


def test_is_supported():
    assert is_supported("x.PDF") and is_supported("y.md") and not is_supported("z.exe")


# --------------------------------------------------------------------------- #
# ERROR — sandbox containment + typed errors (never a raw exception)
# --------------------------------------------------------------------------- #
def test_absolute_path_escape(root):
    with pytest.raises(DocAccessError):
        load_document(root, "/etc/shadow")


def test_dotdot_escape(root):
    with pytest.raises(DocAccessError):
        load_document(root, "../../../../etc/passwd")


def test_symlink_outside_root_refused(root, tmp_path):
    outside = tmp_path.parent / "outside_secret.txt"
    outside.write_text("TOPSECRET", encoding="utf-8")
    link = root / "link.txt"
    try:
        link.symlink_to(outside)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unsupported here")
    with pytest.raises(DocAccessError):
        load_document(root, "link.txt") # resolve follows → outside → refused


def test_resolve_within_allows_in_root(root):
    p = resolve_within(root, "sub/nested.txt")
    assert p.is_file() and p.read_text() == "nested"


def test_missing_file(root):
    with pytest.raises(DocNotFound):
        load_document(root, "nope.txt")


def test_unsupported_format(root):
    with pytest.raises(UnsupportedFormat):
        load_document(root, "ignore.bin")


def test_oversize(root):
    with pytest.raises(DocTooLarge):
        load_document(root, "a.txt", max_bytes=3)


def test_binary_as_txt_is_load_error(root):
    (root / "bad.txt").write_bytes(b"\xff\xfe\x00\x80not utf8\xff")
    with pytest.raises(DocLoadError):
        load_document(root, "bad.txt")


def test_corrupt_pdf_is_load_error(root):
    (root / "broken.pdf").write_bytes(b"%PDF-1.4 not really a pdf at all")
    with pytest.raises(DocLoadError):
        load_document(root, "broken.pdf")


def test_missing_optional_dep_disables_format(root, monkeypatch):
    # simulate pypdf not installed → .pdf yields UnsupportedFormat, not an import crash
    monkeypatch.setitem(sys.modules, "pypdf", None)
    with pytest.raises(UnsupportedFormat):
        load_document(root, "d.pdf")


def test_list_documents_missing_subdir_returns_empty(root):
    assert list_documents(root, "does-not-exist") == []


def test_list_documents_escaping_subdir_raises_access_error(root):
    # by design + symmetric with load_document: an escaping subdir is a typed DocAccessError
    # (NOT a raw exception, NOT a silent []). The gate also guards this arg at the tool level.
    with pytest.raises(DocAccessError):
        list_documents(root, "../../etc")
    with pytest.raises(DocAccessError):
        list_documents(root, "/etc")


def test_permission_denied_file_is_load_error(root):
    # regression (round-3): a chmod 000 text file → DocLoadError, never raw PermissionError
    import os as _os
    p = root / "locked.txt"
    p.write_text("secret", encoding="utf-8")
    _os.chmod(p, 0o000)
    try:
        if _os.access(p, _os.R_OK): # running as root ignores perms → skip
            pytest.skip("cannot drop read permission (running as root?)")
        with pytest.raises(DocLoadError):
            load_document(root, "locked.txt")
    finally:
        _os.chmod(p, 0o644)


def test_list_documents_perm_denied_subdir_no_crash(root):
    # regression (round-3): non-recursive list of a 000 dir → [] (or partial), never raises
    import os as _os
    sub = root / "locked_dir"
    sub.mkdir()
    (sub / "x.txt").write_text("x", encoding="utf-8")
    _os.chmod(sub, 0o000)
    try:
        if _os.access(sub, _os.R_OK):
            pytest.skip("cannot drop read permission (running as root?)")
        assert list_documents(root, "locked_dir", recursive=False) == [] # no crash
    finally:
        _os.chmod(sub, 0o755)


def test_list_documents_excludes_symlink_to_outside_file(root, tmp_path):
    # regression: a symlink inside root → outside file must be EXCLUDED, never crash
    outside = tmp_path.parent / "outside_listed.txt"
    outside.write_text("OUTSIDE", encoding="utf-8")
    try:
        (root / "evil.txt").symlink_to(outside)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unsupported here")
    docs = list_documents(root) # must not raise
    assert "evil.txt" not in docs # outside-pointing symlink excluded
    assert "a.txt" in docs # legit files still listed


def test_list_documents_symlink_loop_does_not_crash(root):
    # regression (round-2): a symlink loop inside root must be excluded, not crash
    try:
        (root / "loop_a.txt").symlink_to(root / "loop_b.txt")
        (root / "loop_b.txt").symlink_to(root / "loop_a.txt")
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unsupported here")
    docs = list_documents(root) # must not raise (RuntimeError ELOOP)
    assert "a.txt" in docs # legit files still listed
    assert "loop_a.txt" not in docs and "loop_b.txt" not in docs


def test_list_documents_does_not_descend_symlinked_dir(root, tmp_path):
    # a symlinked directory inside root → outside dir must not be traversed/listed
    outside_dir = tmp_path.parent / "outside_dir"
    outside_dir.mkdir(exist_ok=True)
    (outside_dir / "leak.txt").write_text("LEAK", encoding="utf-8")
    try:
        (root / "linkdir").symlink_to(outside_dir, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("symlinks unsupported here")
    docs = list_documents(root) # must not raise, must not leak
    assert not any("leak.txt" in d for d in docs)
    assert not any("linkdir" in d for d in docs)


# --------------------------------------------------------------------------- #
# load_bytes — in-memory, format-aware
# --------------------------------------------------------------------------- #
def test_load_bytes_txt_md():
    assert load_bytes("a.txt", b"hello bytes") == "hello bytes"
    assert "body" in load_bytes("b.md", b"# t\n\nbody")


def test_load_bytes_pdf():
    assert "Hello PDF" in load_bytes("d.pdf", _make_pdf("Hello PDF"))


def test_load_bytes_docx():
    import io as _io
    import docx
    d = docx.Document()
    d.add_paragraph("docx bytes line")
    buf = _io.BytesIO()
    d.save(buf)
    assert "docx bytes line" in load_bytes("e.docx", buf.getvalue())


def test_load_bytes_html_strips():
    t = load_bytes("c.html", b"<html><body><h1>Hi</h1><script>evil()</script><p>world</p></body></html>")
    assert "Hi" in t and "world" in t and "evil" not in t


def test_load_bytes_unsupported():
    with pytest.raises(UnsupportedFormat):
        load_bytes("x.bin", b"\x00\x01")


def test_load_bytes_oversize():
    with pytest.raises(DocTooLarge):
        load_bytes("a.txt", b"hello", max_bytes=2)


def test_load_bytes_bad_utf8():
    with pytest.raises(DocLoadError):
        load_bytes("a.txt", b"\xff\xfe not utf8")


def test_load_bytes_corrupt_pdf():
    with pytest.raises(DocLoadError):
        load_bytes("x.pdf", b"this is not a pdf")


def test_load_bytes_non_bytes():
    with pytest.raises(DocLoadError):
        load_bytes("a.txt", "i am a string") # type: ignore[arg-type]
