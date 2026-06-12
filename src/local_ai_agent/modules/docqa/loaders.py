"""Document loaders + the `docs_root` sandbox.

Turn a local file into plain text, *safely*. Every read is contained to a configured
`docs_root` via `resolve_within` (resolve — following symlinks — then `is_relative_to`),
so `..`/absolute/symlink escapes and secret/system paths are refused before any I/O. This
is defense-in-depth: it holds independently of (and on top of) the safety gate.

Loaders are extension-dispatched and lazy-import their optional dependency, so a missing
dep disables only that format (typed error), never an import crash. No loader ever raises a
raw library/OS exception — everything maps to a `DocError`.
"""
from __future__ import annotations

import io
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

_TEXT_EXT = {".txt", ".md", ".markdown", ".text"}
_HTML_EXT = {".html", ".htm"}
# NOTE: `.hwpx` (open OWPML zip) is supported; legacy binary `.hwp` is deliberately NOT —
# it falls through to `UnsupportedFormat`.
_SUPPORTED = _TEXT_EXT | _HTML_EXT | {".pdf", ".docx", ".hwpx"}

_DEFAULT_MAX_BYTES = 10_000_000

# HWPX (zip) bounds — defense against a zip-bomb / over-large member set. The body text
# lives in `Contents/sectionN.xml`; we read ONLY those members, never trust the declared
# size, and bound the decompressed bytes by `_DEFAULT_MAX_BYTES` (a real bounded read, not
# a header check). XML is parsed with stdlib ElementTree, which does not resolve external
# entities (no XXE / no network or file fetch).
_HWPX_MAX_MEMBERS = 4096
_HWPX_SECTION_RE = re.compile(r"(?:^|/)Contents/section(\d+)\.xml$")


# --------------------------------------------------------------------------- #
# errors — every failure is a clean typed DocError (never a raw exception)
# --------------------------------------------------------------------------- #
class DocError(Exception):
    """Base for all document-loading failures."""


class DocAccessError(DocError):
    """The path escapes `docs_root` (absolute / `..` / symlink) — refused before any read."""


class DocNotFound(DocError):
    """No such file under `docs_root`."""


class UnsupportedFormat(DocError):
    """Extension not supported (or its optional dependency is unavailable)."""


class DocTooLarge(DocError):
    """File exceeds the configured byte cap."""


class DocLoadError(DocError):
    """The file could not be parsed/decoded (malformed, binary, corrupt)."""


# --------------------------------------------------------------------------- #
# the sandbox
# --------------------------------------------------------------------------- #
def resolve_within(root: Path, relpath: str | Path) -> Path:
    """Resolve `relpath` under `root` and require containment. `resolve()` runs first so a
    symlink is followed *then* checked — a symlink pointing outside `root` is rejected. An
    absolute `relpath` (which resets the join) or a `..` escape also resolves outside →
    `DocAccessError`."""
    root_resolved = Path(root).resolve()
    try:
        target = (root_resolved / Path(relpath)).resolve()
    except (OSError, RuntimeError, ValueError) as exc: # e.g. symlink loop / bad path
        raise DocAccessError(f"cannot resolve path: {relpath!r}") from exc
    if target != root_resolved and root_resolved not in target.parents:
        raise DocAccessError(f"path escapes docs_root: {relpath!r}")
    return target


# --------------------------------------------------------------------------- #
# per-format loaders (lazy imports → a missing dep disables just that format)
# --------------------------------------------------------------------------- #
def _load_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise DocLoadError(f"not valid UTF-8 text: {path.name}") from exc
    except OSError as exc: # permission denied, I/O error, etc. → typed, never raw
        raise DocLoadError(f"could not read: {path.name}") from exc


def _load_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UnsupportedFormat("pdf support not installed (pypdf)") from exc
    try:
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc: # noqa: BLE001 — any pypdf failure → typed error, never raw
        raise DocLoadError(f"could not parse PDF: {path.name}") from exc


def _load_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedFormat("docx support not installed (python-docx)") from exc
    try:
        document = docx.Document(str(path))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception as exc: # noqa: BLE001
        raise DocLoadError(f"could not parse DOCX: {path.name}") from exc


def _load_html(path: Path) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise UnsupportedFormat("html support not installed (beautifulsoup4)") from exc
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text(separator=" ", strip=True)
    except Exception as exc: # noqa: BLE001
        raise DocLoadError(f"could not parse HTML: {path.name}") from exc


# --------------------------------------------------------------------------- #
# HWPX (open OWPML zip) — stdlib zip + XML, bounded (zip-bomb) and XXE-off.
# Shared core used by both the path and bytes loaders.
# --------------------------------------------------------------------------- #
def _hwpx_local(tag: object) -> str:
    """The namespace-stripped local name of an element tag (HWPX namespaces vary)."""
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _hwpx_section_text(raw: bytes) -> str:
    """Plain text of one `section*.xml`: join `<*:t>` runs per `<*:p>` paragraph. ElementTree
    does not resolve external entities (no XXE); a malformed/entity-bearing doc → ParseError →
    typed `DocLoadError` (no fetch, no crash)."""
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise DocLoadError("could not parse HWPX section XML") from exc
    paragraphs: list[str] = []
    for para in root.iter():
        if _hwpx_local(para.tag) != "p":
            continue
        runs = [t.text for t in para.iter() if _hwpx_local(t.tag) == "t" and t.text]
        if runs:
            paragraphs.append("".join(runs))
    if not paragraphs: # fallback: any text run anywhere in the section
        flat = "".join(t.text for t in root.iter() if _hwpx_local(t.tag) == "t" and t.text)
        if flat:
            paragraphs.append(flat)
    return "\n".join(paragraphs)


def _hwpx_text_from_zip(zf: zipfile.ZipFile, max_bytes: int) -> str:
    """Extract body text from `Contents/sectionN.xml` members, in section order. Bounded:
    member-count cap + a real bounded read (never decompress past `max_bytes`, the declared
    size is never trusted)."""
    if len(zf.infolist()) > _HWPX_MAX_MEMBERS:
        raise DocTooLarge(f"hwpx has too many members (> {_HWPX_MAX_MEMBERS})")
    sections: list[tuple[int, str]] = []
    for name in zf.namelist():
        m = _HWPX_SECTION_RE.search(name)
        if m:
            sections.append((int(m.group(1)), name))
    if not sections:
        raise DocLoadError("hwpx has no Contents/section*.xml")
    sections.sort()
    texts: list[str] = []
    remaining = max_bytes
    for _, name in sections:
        with zf.open(name) as fh: # bounded read: +1 to detect an overflow
            raw = fh.read(remaining + 1)
        if len(raw) > remaining:
            raise DocTooLarge(f"hwpx content exceeds {max_bytes} bytes")
        remaining -= len(raw)
        texts.append(_hwpx_section_text(raw))
    return "\n".join(t for t in texts if t).strip()


def _load_hwpx(path: Path) -> str:
    try:
        with zipfile.ZipFile(str(path)) as zf:
            return _hwpx_text_from_zip(zf, _DEFAULT_MAX_BYTES)
    except zipfile.BadZipFile as exc:
        raise DocLoadError(f"not a valid HWPX (zip): {path.name}") from exc
    except DocError:
        raise
    except Exception as exc: # noqa: BLE001 — any zip/XML failure → typed error, never raw
        raise DocLoadError(f"could not parse HWPX: {path.name}") from exc


_LOADERS = {
    **{ext: _load_text for ext in _TEXT_EXT},
    **{ext: _load_html for ext in _HTML_EXT},
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".hwpx": _load_hwpx,
}


def is_supported(name: str | Path) -> bool:
    return Path(name).suffix.lower() in _SUPPORTED


# --------------------------------------------------------------------------- #
# public API
# --------------------------------------------------------------------------- #
def load_document(root: Path, relpath: str | Path, *, max_bytes: int = _DEFAULT_MAX_BYTES) -> str:
    """Load a contained document as plain text. Order: contain → exists/is_file → size →
    dispatch by extension. Every failure is a typed `DocError`."""
    target = resolve_within(root, relpath) # raises DocAccessError on any escape (guarded)
    # backstop: after containment, ANY raw OS/lib exception maps to a typed DocError — the
    # "never raise raw" invariant must hold for every filesystem quirk (perms, ELOOP, FIFO…).
    try:
        if not target.is_file():
            raise DocNotFound(f"no such document: {relpath!r}")
        size = target.stat().st_size
        if size > max_bytes:
            raise DocTooLarge(f"document exceeds {max_bytes} bytes: {relpath!r}")
        loader = _LOADERS.get(target.suffix.lower())
        if loader is None:
            raise UnsupportedFormat(f"unsupported format: {target.suffix!r}")
        return loader(target)
    except DocError:
        raise
    except Exception as exc: # noqa: BLE001 — robustness backstop; never propagate a raw error
        raise DocLoadError(f"could not load {relpath!r}: {type(exc).__name__}") from exc


# --------------------------------------------------------------------------- #
# in-memory (bytes) loading — for content fetched from a non-local source
# (e.g. a remote storage connector, ). Mirrors `load_document`'s
# extension dispatch + typed-error discipline, parsing from a buffer (no path).
# --------------------------------------------------------------------------- #
def _load_text_bytes(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise DocLoadError("not valid UTF-8 text") from exc


def _load_pdf_bytes(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UnsupportedFormat("pdf support not installed (pypdf)") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc: # noqa: BLE001
        raise DocLoadError("could not parse PDF") from exc


def _load_docx_bytes(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise UnsupportedFormat("docx support not installed (python-docx)") from exc
    try:
        document = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    except Exception as exc: # noqa: BLE001
        raise DocLoadError("could not parse DOCX") from exc


def _load_html_bytes(data: bytes) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:
        raise UnsupportedFormat("html support not installed (beautifulsoup4)") from exc
    try:
        soup = BeautifulSoup(data.decode("utf-8", "replace"), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text(separator=" ", strip=True)
    except Exception as exc: # noqa: BLE001
        raise DocLoadError("could not parse HTML") from exc


def _load_hwpx_bytes(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return _hwpx_text_from_zip(zf, _DEFAULT_MAX_BYTES)
    except zipfile.BadZipFile as exc:
        raise DocLoadError("not a valid HWPX (zip)") from exc
    except DocError:
        raise
    except Exception as exc: # noqa: BLE001
        raise DocLoadError("could not parse HWPX") from exc


_BYTES_LOADERS = {
    **{ext: _load_text_bytes for ext in _TEXT_EXT},
    **{ext: _load_html_bytes for ext in _HTML_EXT},
    ".pdf": _load_pdf_bytes,
    ".docx": _load_docx_bytes,
    ".hwpx": _load_hwpx_bytes,
}


def load_bytes(filename: str | Path, data: bytes, *, max_bytes: int = _DEFAULT_MAX_BYTES) -> str:
    """Load already-fetched `data` as plain text, dispatching by `filename`'s extension (no
    filesystem access). Same typed-error discipline as `load_document`: oversize → `DocTooLarge`,
    unknown ext → `UnsupportedFormat`, any parse failure → `DocLoadError` (never a raw exception)."""
    if not isinstance(data, (bytes, bytearray)):
        raise DocLoadError("data must be bytes")
    if len(data) > max_bytes:
        raise DocTooLarge(f"document exceeds {max_bytes} bytes")
    loader = _BYTES_LOADERS.get(Path(filename).suffix.lower())
    if loader is None:
        raise UnsupportedFormat(f"unsupported format: {Path(filename).suffix!r}")
    try:
        return loader(bytes(data))
    except DocError:
        raise
    except Exception as exc: # noqa: BLE001 — robustness backstop; never propagate a raw error
        raise DocLoadError(f"could not load {str(filename)!r}: {type(exc).__name__}") from exc


def _relative_within(path: Path, root_resolved: Path) -> str | None:
    """Resolved-root-relative string, or None if `path` resolves outside root (e.g. a
    symlink pointing out) or can't be resolved — such entries are silently excluded, never
    raising. Closes the `list_documents` symlink-escape crash."""
    try:
        rel = path.resolve().relative_to(root_resolved)
    except (ValueError, OSError, RuntimeError): # RuntimeError: py3.11 symlink-loop (ELOOP)
        return None
    return str(rel)


def list_documents(root: Path, subdir: str | Path = ".", *, recursive: bool = True) -> list[str]:
    """Sorted, root-relative paths of supported, non-hidden files under `root`/`subdir`
    (contained). Deterministic. Uses `os.walk(followlinks=False)` so symlinked directories
    are not descended, and a per-file containment check excludes any file-symlink that
    resolves outside root — never raising on such entries."""
    base = resolve_within(root, subdir)
    root_resolved = Path(root).resolve()
    if not base.is_dir():
        return []
    out: list[str] = []
    try:
        if recursive:
            for dirpath, dirnames, filenames in os.walk(base, followlinks=False):
                dirnames[:] = [d for d in dirnames if not d.startswith(".")] # prune hidden dirs
                for fn in filenames:
                    if fn.startswith(".") or not is_supported(fn):
                        continue
                    rel = _relative_within(Path(dirpath) / fn, root_resolved)
                    if rel is not None:
                        out.append(rel)
        else:
            for p in base.iterdir():
                if not p.is_file() or p.name.startswith(".") or not is_supported(p):
                    continue
                rel = _relative_within(p, root_resolved)
                if rel is not None:
                    out.append(rel)
    except (OSError, RuntimeError, ValueError): # any listing-time FS error → partial, never raise
        pass
    return sorted(out)
